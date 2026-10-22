import io
import docx
from docx.shared import Pt, Inches
from reportlab.lib.pagesizes import letter, A4
from reportlab.pdfgen import canvas

from app.services.compliance.university_profiles import get_profile, PROFILES
from app.services.compliance.compliance_checker import DocumentComplianceAnalyzer
from app.services.compliance.docx_smart_fixer import DocxSmartFixer
from app.models.compliance import IssueCategory, IssueSeverity


def test_university_profiles():
    for uid in ["INSAT", "ESPRIT", "ENSI", "FST", "ISET"]:
        prof = get_profile(uid)
        assert prof.id == uid
        assert len(prof.body_font_families) > 0
        assert prof.line_spacing == 1.5
        assert "top" in prof.margins_cm
        assert prof.figure_caption_pattern
        assert prof.table_caption_pattern
    print("[OK] University profiles verified")


def test_docx_smart_fixer():
    doc = docx.Document()
    # Add an unformatted paragraph
    p = doc.add_paragraph("Ceci est un paragraphe de test qui n'est pas aux normes.")
    p.paragraph_format.line_spacing = 1.0
    
    # Add an unformatted figure caption
    p_fig = doc.add_paragraph("Fig. 1 diagramme d'architecture")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    
    buf = io.BytesIO()
    doc.save(buf)
    raw_docx = buf.getvalue()
    
    profile = get_profile("INSAT")
    fixer = DocxSmartFixer(raw_docx, profile)
    fixed_bytes, fixes = fixer.fix()
    
    assert len(fixed_bytes) > 0
    assert len(fixes) >= 3
    
    # Verify fixed docx
    fixed_doc = docx.Document(io.BytesIO(fixed_bytes))
    # Margins: INSAT has 2.5cm
    for sec in fixed_doc.sections:
        assert round(sec.top_margin.cm, 1) == 2.5
        assert round(sec.left_margin.cm, 1) == 2.5
    
    print("[OK] DocxSmartFixer tested successfully:", fixes[0])


def test_pdf_compliance_analyzer():
    # Build a synthetic PDF with ReportLab
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    
    # Page 1: Cover
    c.setFont("Helvetica-Bold", 18)
    c.drawString(100, 750, "Projet de Fin d'Etudes - PFE")
    c.setFont("Helvetica", 12)
    c.drawString(100, 700, "Université de Carthage - INSAT")
    c.showPage()
    
    # Page 2: Table des matières
    c.setFont("Helvetica-Bold", 16)
    c.drawString(100, 750, "Table des matières")
    c.setFont("Helvetica", 11)
    c.drawString(100, 700, "1. Introduction générale ................................ 3")
    c.drawString(100, 680, "2. Analyse et Conception ................................. 4")
    c.showPage()
    
    # Page 3: Introduction
    c.setFont("Helvetica-Bold", 16)
    c.drawString(100, 750, "1. Introduction générale")
    c.setFont("Helvetica", 12)
    c.drawString(100, 700, "Dans le cadre de notre projet de fin d'études, nous développons un système.")
    c.drawString(100, 680, "Ce paragraphe illustre l'alignement et la typographie requise.")
    c.drawString(100, 600, "Figure 1.1: Architecture globale de la plateforme")
    c.showPage()
    
    c.save()
    pdf_bytes = buf.getvalue()
    
    profile = get_profile("INSAT")
    analyzer = DocumentComplianceAnalyzer(pdf_bytes, profile, file_name="pfe_insat.pdf")
    results = analyzer.analyze()
    
    assert "overall_score" in results
    assert "compliance_level" in results
    assert "category_scores" in results
    assert "issues" in results
    assert results["stats"]["page_count"] == 3
    assert results["stats"]["detected_figures"] >= 1
    assert results["stats"]["toc_entries_count"] >= 1
    print(f"[OK] PDF Analyzer score: {results['overall_score']}% ({results['compliance_level']})")


if __name__ == "__main__":
    test_university_profiles()
    test_docx_smart_fixer()
    test_pdf_compliance_analyzer()
    print("ALL COMPLIANCE TESTS PASSED SUCCESSFULLY!")
